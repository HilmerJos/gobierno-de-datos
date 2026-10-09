import os
import json
import shutil
from datetime import datetime
from typing import Optional, List
from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
import openpyxl
import docx
from sqlalchemy import create_engine, Column, Integer, String, Float, Text, Date, DateTime, text
from sqlalchemy.orm import declarative_base, sessionmaker, Session

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(BASE_DIR)
UPLOADS_DIR = os.path.join(PROJECT_ROOT, "uploads")
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
FRONTEND_DIST = os.path.join(PROJECT_ROOT, "frontend", "dist")
STATIC_DIR = os.path.join(PROJECT_ROOT, "static")
CATALOG_PATH = os.path.join(BASE_DIR, "catalogo_preguntas_pcm.json")
SQLITE_PATH = os.path.join(BASE_DIR, "sgd_database.db")

os.makedirs(UPLOADS_DIR, exist_ok=True)
os.makedirs(TEMPLATES_DIR, exist_ok=True)

# -------------------------------------------------------------
# CONFIGURACIÓN DE BASE DE DATOS (POSTGRESQL CON FALLBACK A SQLITE)
# -------------------------------------------------------------
PG_URL = os.getenv("DATABASE_URL", "postgresql+psycopg2://postgres:postgres@localhost:5432/sgd_madurez")
DB_ENGINE_NAME = "sqlite"

try:
    # Intento de conexión con PostgreSQL
    pg_engine = create_engine(PG_URL, pool_pre_ping=True, connect_args={"connect_timeout": 3})
    with pg_engine.connect() as test_conn:
        test_conn.execute(text("SELECT 1"))
    engine = pg_engine
    DB_ENGINE_NAME = "postgresql"
    print(f"[SGD-Madurez] Conectado exitosamente a PostgreSQL: {PG_URL}")
except Exception as e:
    # Fallback transparente a SQLite para garantizar operatividad continua
    sqlite_url = f"sqlite:///{SQLITE_PATH}"
    engine = create_engine(sqlite_url, connect_args={"check_same_thread": False})
    DB_ENGINE_NAME = "sqlite"
    print(f"[SGD-Madurez] Usando base de datos SQLite (fallback): {sqlite_url}")

SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

# -------------------------------------------------------------
# MODELOS ORM
# -------------------------------------------------------------
class DimensionModel(Base):
    __tablename__ = "dimensiones"
    id = Column(Integer, primary_key=True, index=True)
    codigo = Column(String(10), unique=True, index=True)
    nombre = Column(String(255))
    orden = Column(Integer)

class PreguntaModel(Base):
    __tablename__ = "preguntas"
    id = Column(Integer, primary_key=True, index=True)
    dimension_codigo = Column(String(10), index=True)
    codigo = Column(String(20), unique=True, index=True)
    enunciado = Column(Text)
    explicacion = Column(Text)
    resultados_esperados = Column(Text)
    orden = Column(Integer)

class OpcionModel(Base):
    __tablename__ = "opciones"
    id = Column(Integer, primary_key=True, autoincrement=True)
    pregunta_codigo = Column(String(20), index=True)
    nivel = Column(Integer)
    etiqueta = Column(String(100))
    descripcion = Column(Text)

class RespuestaModel(Base):
    __tablename__ = "respuestas"
    pregunta_codigo = Column(String(20), primary_key=True)
    valor_nivel = Column(Float)
    opcion_texto = Column(Text)
    comentarios = Column(Text, default="")
    actualizado_en = Column(String(50))

class EvidenciaModel(Base):
    __tablename__ = "evidencias"
    id = Column(Integer, primary_key=True, autoincrement=True)
    pregunta_codigo = Column(String(20), index=True)
    origen = Column(String(100))
    tipo_evidencia = Column(String(100))
    nombre_evidencia = Column(String(255))
    ruta_o_enlace = Column(Text)
    descripcion = Column(Text)
    responsable = Column(String(200))
    area_unidad = Column(String(200))
    fecha_registro = Column(String(50))
    estado_validacion = Column(String(50), default="En Revisión")
    comentarios_revision = Column(Text, default="")

class ConfigInstitucionalModel(Base):
    __tablename__ = "configuracion_institucional"
    id = Column(Integer, primary_key=True)
    entidad_nombre = Column(String(255))
    siglas = Column(String(50))
    ogd_nombre = Column(String(200))
    ogd_puesto = Column(String(200))
    presidente_cgtd = Column(String(200))
    resolucion_cgtd = Column(String(100))
    periodo = Column(String(50))
    nivel_objetivo = Column(Float, default=3.0)

class IniciativaModel(Base):
    __tablename__ = "iniciativas_plan"
    id = Column(Integer, primary_key=True, autoincrement=True)
    codigo = Column(String(50))
    nombre = Column(String(255))
    descripcion = Column(Text)
    eje_engd = Column(String(150))
    plazo = Column(String(50))
    peso_relativo = Column(Float, default=10.0)
    meta_porcentaje = Column(Float, default=100.0)
    avance_porcentaje = Column(Float, default=0.0)
    responsable = Column(String(200))
    puesto = Column(String(200))
    unidad_organica = Column(String(200))
    costo_proyectado = Column(Float, default=0.0)
    costo_ejecutado = Column(Float, default=0.0)
    fuente_financiamiento = Column(String(100))
    fecha_inicio = Column(String(50))
    fecha_fin = Column(String(50))

# Crear tablas
Base.metadata.create_all(bind=engine)

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# -------------------------------------------------------------
# SEEDER AUTOMÁTICO INICIAL
# -------------------------------------------------------------
def seed_database():
    db = SessionLocal()
    try:
        q_count = db.query(PreguntaModel).count()
        if q_count == 0 and os.path.exists(CATALOG_PATH):
            with open(CATALOG_PATH, "r", encoding="utf-8") as f:
                catalog = json.load(f)
            for d in catalog:
                db.merge(DimensionModel(id=d["id"], codigo=d["code"], nombre=d["name"], orden=d["id"]))
                for q in d["questions"]:
                    db.merge(PreguntaModel(
                        id=q["id"],
                        dimension_codigo=d["code"],
                        codigo=q["code"],
                        enunciado=q["enunciado"],
                        explicacion=q["explicacion"],
                        resultados_esperados=q["resultados_esperados"],
                        orden=q["id"]
                    ))
                    for opt in q["options"]:
                        db.add(OpcionModel(
                            pregunta_codigo=q["code"],
                            nivel=opt["nivel"],
                            etiqueta=opt["etiqueta"],
                            descripcion=opt["descripcion"]
                        ))
            db.commit()

        cfg_count = db.query(ConfigInstitucionalModel).count()
        if cfg_count == 0:
            db.add(ConfigInstitucionalModel(
                id=1,
                entidad_nombre="Oficina de Tecnologías de la Información",
                siglas="OTI",
                ogd_nombre="Ing. Oficial de Gobierno de Datos",
                ogd_puesto="Jefe de OTI / Oficial de Datos",
                presidente_cgtd="Director General",
                resolucion_cgtd="Resolución N° 001-2026-CGTD",
                periodo="2026 - 2030",
                nivel_objetivo=3.0
            ))
            db.commit()

        init_count = db.query(IniciativaModel).count()
        if init_count == 0:
            sample_inits = [
                IniciativaModel(
                    codigo="INIC-01",
                    nombre="Aprobación y Difusión de la Estrategia de Gobierno de Datos (EGD)",
                    descripcion="Formalizar la EGD institucional y su incorporación en el Plan de Gobierno Digital.",
                    eje_engd="Eje 1: Gobernanza e Institucionalidad",
                    plazo="Corto plazo",
                    peso_relativo=20.0,
                    meta_porcentaje=100.0,
                    avance_porcentaje=35.0,
                    responsable="Oficial de Gobierno de Datos",
                    puesto="Jefe OTI",
                    unidad_organica="OTI",
                    costo_proyectado=15000.0,
                    costo_ejecutado=5000.0,
                    fuente_financiamiento="Recursos Ordinarios",
                    fecha_inicio="2026-03-01",
                    fecha_fin="2026-08-31"
                ),
                IniciativaModel(
                    codigo="INIC-02",
                    nombre="Implementación del Modelo de Referencia y Glosario de Datos",
                    descripcion="Construir el inventario de datos maestros y glosario conceptual centralizado.",
                    eje_engd="Eje 3: Calidad e Interoperabilidad de Datos",
                    plazo="Mediano plazo",
                    peso_relativo=25.0,
                    meta_porcentaje=100.0,
                    avance_porcentaje=20.0,
                    responsable="Arquitecto de Datos",
                    puesto="Especialista BD",
                    unidad_organica="Área de Desarrollo y BD",
                    costo_proyectado=28000.0,
                    costo_ejecutado=0.0,
                    fuente_financiamiento="Recursos Ordinarios",
                    fecha_inicio="2026-05-01",
                    fecha_fin="2026-12-31"
                ),
                IniciativaModel(
                    codigo="INIC-03",
                    nombre="Programa Institucional de Calidad de Datos en Sistemas Críticos",
                    descripcion="Definir métricas, umbrales y reglas de negocio para depuración y monitoreo de datos.",
                    eje_engd="Eje 3: Calidad e Interoperabilidad de Datos",
                    plazo="Mediano plazo",
                    peso_relativo=20.0,
                    meta_porcentaje=100.0,
                    avance_porcentaje=10.0,
                    responsable="Especialista de Calidad",
                    puesto="Analista de Datos",
                    unidad_organica="OTI",
                    costo_proyectado=20000.0,
                    costo_ejecutado=0.0,
                    fuente_financiamiento="Recursos Ordinarios",
                    fecha_inicio="2026-07-01",
                    fecha_fin="2027-04-30"
                ),
                IniciativaModel(
                    codigo="INIC-04",
                    nombre="Publicación Sistemática de Datasets Abiertos en la PNDA",
                    descripcion="Identificar y anonimizar datos prioritarios de alto valor público para la Plataforma Nacional de Datos Abiertos.",
                    eje_engd="Eje 5: Uso Estratégico y Ético de los Datos",
                    plazo="Corto plazo",
                    peso_relativo=15.0,
                    meta_porcentaje=100.0,
                    avance_porcentaje=50.0,
                    responsable="Coordinador de Interoperabilidad",
                    puesto="Especialista PIDE",
                    unidad_organica="OTI",
                    costo_proyectado=8000.0,
                    costo_ejecutado=4000.0,
                    fuente_financiamiento="Recursos Ordinarios",
                    fecha_inicio="2026-04-01",
                    fecha_fin="2026-10-31"
                ),
                IniciativaModel(
                    codigo="INIC-05",
                    nombre="Programa de Alfabetización y Cultura de Datos (Data Literacy)",
                    descripcion="Capacitaciones continuas a los tomadores de decisiones y custodios sobre el valor del dato.",
                    eje_engd="Eje 4: Cultura y Capacidades en Datos",
                    plazo="Largo plazo",
                    peso_relativo=20.0,
                    meta_porcentaje=100.0,
                    avance_porcentaje=15.0,
                    responsable="Especialista en Gestión del Cambio",
                    puesto="Capacitador RRHH",
                    unidad_organica="Recursos Humanos y OTI",
                    costo_proyectado=12000.0,
                    costo_ejecutado=2000.0,
                    fuente_financiamiento="Recursos Ordinarios",
                    fecha_inicio="2026-06-01",
                    fecha_fin="2027-11-30"
                )
            ]
            for init in sample_inits:
                db.add(init)
            db.commit()
    finally:
        db.close()

seed_database()

# -------------------------------------------------------------
# APLICACIÓN FASTAPI
# -------------------------------------------------------------
app = FastAPI(
    title="SGD-Madurez API",
    description="Sistema de Evaluación de Madurez de Gobierno de Datos y Gestión del Plan de Acción",
    version="2.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# -------------------------------------------------------------
# MODELOS PYDANTIC
# -------------------------------------------------------------
class RespuestaIn(BaseModel):
    pregunta_codigo: str
    valor_nivel: float
    opcion_texto: str
    comentarios: Optional[str] = ""

class EvidenciaIn(BaseModel):
    pregunta_codigo: str
    origen: str
    tipo_evidencia: str
    nombre_evidencia: str
    ruta_o_enlace: str
    descripcion: str
    responsable: str
    area_unidad: str
    fecha_registro: Optional[str] = None
    estado_validacion: Optional[str] = "En Revisión"
    comentarios_revision: Optional[str] = ""

class ConfigInstitucionalIn(BaseModel):
    entidad_nombre: str
    siglas: str
    ogd_nombre: str
    ogd_puesto: str
    presidente_cgtd: str
    resolucion_cgtd: str
    periodo: str
    nivel_objetivo: float

class IniciativaIn(BaseModel):
    codigo: str
    nombre: str
    descripcion: str
    eje_engd: str
    plazo: str
    peso_relativo: float
    meta_porcentaje: float
    avance_porcentaje: float
    responsable: str
    puesto: str
    unidad_organica: str
    costo_proyectado: float
    costo_ejecutado: float
    fuente_financiamiento: str
    fecha_inicio: str
    fecha_fin: str

# -------------------------------------------------------------
# ENDPOINTS REST
# -------------------------------------------------------------
@app.get("/api/db-info")
def db_info():
    return {
        "engine": DB_ENGINE_NAME,
        "status": "connected",
        "timestamp": datetime.now().isoformat()
    }

@app.get("/api/dimensiones")
def listar_dimensiones(db: Session = Depends(get_db)):
    dims = db.query(DimensionModel).order_by(DimensionModel.orden).all()
    result = []
    for d in dims:
        pregs = db.query(PreguntaModel).filter(PreguntaModel.dimension_codigo == d.codigo).all()
        total_p = len(pregs)
        preg_codes = [p.codigo for p in pregs]
        
        resp_rows = db.query(RespuestaModel).filter(RespuestaModel.pregunta_codigo.in_(preg_codes)).all()
        respondidas = len(resp_rows)
        score = sum(r.valor_nivel for r in resp_rows) / respondidas if respondidas > 0 else 0.0
        
        result.append({
            "id": d.id,
            "codigo": d.codigo,
            "nombre": d.nombre,
            "total_preguntas": total_p,
            "respondidas": respondidas,
            "puntaje": round(score, 2)
        })
    return result

@app.get("/api/preguntas")
def listar_preguntas(dim: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(PreguntaModel)
    if dim:
        query = query.filter(PreguntaModel.dimension_codigo == dim)
    pregs = query.order_by(PreguntaModel.id).all()
    
    result = []
    for p in pregs:
        opts = db.query(OpcionModel).filter(OpcionModel.pregunta_codigo == p.codigo).order_by(OpcionModel.nivel).all()
        ans = db.query(RespuestaModel).filter(RespuestaModel.pregunta_codigo == p.codigo).first()
        ev_count = db.query(EvidenciaModel).filter(EvidenciaModel.pregunta_codigo == p.codigo).count()
        
        result.append({
            "id": p.id,
            "dimension_codigo": p.dimension_codigo,
            "codigo": p.codigo,
            "enunciado": p.enunciado,
            "explicacion": p.explicacion,
            "resultados_esperados": p.resultados_esperados,
            "opciones": [{
                "id": o.id,
                "nivel": o.nivel,
                "etiqueta": o.etiqueta,
                "descripcion": o.descripcion
            } for o in opts],
            "respuesta": {
                "valor_nivel": ans.valor_nivel,
                "opcion_texto": ans.opcion_texto,
                "comentarios": ans.comentarios,
                "actualizado_en": ans.actualizado_en
            } if ans else None,
            "total_evidencias": ev_count
        })
    return result

@app.post("/api/respuestas")
def guardar_respuesta(data: RespuestaIn, db: Session = Depends(get_db)):
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    existente = db.query(RespuestaModel).filter(RespuestaModel.pregunta_codigo == data.pregunta_codigo).first()
    if existente:
        existente.valor_nivel = data.valor_nivel
        existente.opcion_texto = data.opcion_texto
        existente.comentarios = data.comentarios
        existente.actualizado_en = ahora
    else:
        db.add(RespuestaModel(
            pregunta_codigo=data.pregunta_codigo,
            valor_nivel=data.valor_nivel,
            opcion_texto=data.opcion_texto,
            comentarios=data.comentarios,
            actualizado_en=ahora
        ))
    db.commit()
    return {"status": "ok", "message": f"Respuesta guardada para {data.pregunta_codigo}"}

@app.get("/api/estadisticas")
def obtener_estadisticas(db: Session = Depends(get_db)):
    dims = db.query(DimensionModel).order_by(DimensionModel.orden).all()
    dim_scores = []
    sum_dim_scores = 0.0

    for d in dims:
        pregs = db.query(PreguntaModel).filter(PreguntaModel.dimension_codigo == d.codigo).all()
        total_p = len(pregs)
        preg_codes = [p.codigo for p in pregs]
        resp_rows = db.query(RespuestaModel).filter(RespuestaModel.pregunta_codigo.in_(preg_codes)).all()
        respondidas = len(resp_rows)
        score = sum(r.valor_nivel for r in resp_rows) / respondidas if respondidas > 0 else 0.0
        sum_dim_scores += score

        dim_scores.append({
            "codigo": d.codigo,
            "nombre": d.nombre,
            "total_preguntas": total_p,
            "respondidas": respondidas,
            "puntaje": round(score, 2),
            "brecha": round(max(0.0, 3.0 - score), 2)
        })

    puntaje_final = round(sum_dim_scores / len(dims), 2) if len(dims) > 0 else 0.0

    def nivel_verbal(val):
        if val < 0.5: return "Nivel 0: No Iniciado"
        if val < 1.5: return "Nivel 1: Inicial / Ad-hoc"
        if val < 2.5: return "Nivel 2: En Progreso / Parcial"
        if val < 3.5: return "Nivel 3: Definido / Implementado"
        if val < 4.5: return "Nivel 4: Gestionado / Monitoreado"
        return "Nivel 5: Optimizado / Integrado"

    total_preguntas = db.query(PreguntaModel).count()
    total_respondidas = db.query(RespuestaModel).count()
    total_evidencias = db.query(EvidenciaModel).count()
    total_validadas = db.query(EvidenciaModel).filter(EvidenciaModel.estado_validacion == "Validado").count()
    cfg = db.query(ConfigInstitucionalModel).filter(ConfigInstitucionalModel.id == 1).first()
    nivel_obj = cfg.nivel_objetivo if cfg else 3.0

    return {
        "puntaje_final": puntaje_final,
        "nivel_descriptivo": nivel_verbal(puntaje_final),
        "nivel_objetivo": nivel_obj,
        "brecha_global": round(max(0.0, nivel_obj - puntaje_final), 2),
        "total_preguntas": total_preguntas,
        "total_respondidas": total_respondidas,
        "porcentaje_avance": round((total_respondidas / total_preguntas) * 100, 1) if total_preguntas > 0 else 0,
        "total_evidencias": total_evidencias,
        "total_validadas": total_validadas,
        "dimensiones": dim_scores
    }

@app.get("/api/evidencias")
def listar_evidencias(pregunta_codigo: Optional[str] = None, estado: Optional[str] = None, db: Session = Depends(get_db)):
    query = db.query(EvidenciaModel, PreguntaModel.enunciado, PreguntaModel.dimension_codigo)\
        .outerjoin(PreguntaModel, EvidenciaModel.pregunta_codigo == PreguntaModel.codigo)
    
    if pregunta_codigo:
        query = query.filter(EvidenciaModel.pregunta_codigo == pregunta_codigo)
    if estado:
        query = query.filter(EvidenciaModel.estado_validacion == estado)
        
    rows = query.order_by(EvidenciaModel.id.desc()).all()
    res = []
    for ev, preg_enunciado, dim_codigo in rows:
        res.append({
            "id": ev.id,
            "pregunta_codigo": ev.pregunta_codigo,
            "pregunta_enunciado": preg_enunciado,
            "dimension_codigo": dim_codigo,
            "origen": ev.origen,
            "tipo_evidencia": ev.tipo_evidencia,
            "nombre_evidencia": ev.nombre_evidencia,
            "ruta_o_enlace": ev.ruta_o_enlace,
            "descripcion": ev.descripcion,
            "responsable": ev.responsable,
            "area_unidad": ev.area_unidad,
            "fecha_registro": ev.fecha_registro,
            "estado_validacion": ev.estado_validacion,
            "comentarios_revision": ev.comentarios_revision
        })
    return res

@app.post("/api/evidencias")
def crear_evidencia(data: EvidenciaIn, db: Session = Depends(get_db)):
    fecha = data.fecha_registro or datetime.now().strftime("%Y-%m-%d")
    ev = EvidenciaModel(
        pregunta_codigo=data.pregunta_codigo,
        origen=data.origen,
        tipo_evidencia=data.tipo_evidencia,
        nombre_evidencia=data.nombre_evidencia,
        ruta_o_enlace=data.ruta_o_enlace,
        descripcion=data.descripcion,
        responsable=data.responsable,
        area_unidad=data.area_unidad,
        fecha_registro=fecha,
        estado_validacion=data.estado_validacion or "En Revisión",
        comentarios_revision=data.comentarios_revision or ""
    )
    db.add(ev)
    db.commit()
    db.refresh(ev)
    return {"status": "ok", "id": ev.id, "message": "Evidencia registrada exitosamente"}

@app.post("/api/evidencias/upload")
async def subir_archivo_evidencia(file: UploadFile = File(...)):
    filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{file.filename.replace(' ', '_')}"
    filepath = os.path.join(UPLOADS_DIR, filename)
    with open(filepath, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    return {"status": "ok", "filename": filename, "url": f"/uploads/{filename}"}

@app.put("/api/evidencias/{id}")
def actualizar_evidencia(id: int, data: EvidenciaIn, db: Session = Depends(get_db)):
    ev = db.query(EvidenciaModel).filter(EvidenciaModel.id == id).first()
    if not ev:
        raise HTTPException(status_code=404, detail="Evidencia no encontrada")
    ev.pregunta_codigo = data.pregunta_codigo
    ev.origen = data.origen
    ev.tipo_evidencia = data.tipo_evidencia
    ev.nombre_evidencia = data.nombre_evidencia
    ev.ruta_o_enlace = data.ruta_o_enlace
    ev.descripcion = data.descripcion
    ev.responsable = data.responsable
    ev.area_unidad = data.area_unidad
    ev.estado_validacion = data.estado_validacion
    ev.comentarios_revision = data.comentarios_revision
    db.commit()
    return {"status": "ok", "message": "Evidencia actualizada"}

@app.delete("/api/evidencias/{id}")
def eliminar_evidencia(id: int, db: Session = Depends(get_db)):
    ev = db.query(EvidenciaModel).filter(EvidenciaModel.id == id).first()
    if ev:
        db.delete(ev)
        db.commit()
    return {"status": "ok", "message": "Evidencia eliminada"}

@app.get("/api/plan-accion")
def obtener_plan_accion(db: Session = Depends(get_db)):
    cfg = db.query(ConfigInstitucionalModel).filter(ConfigInstitucionalModel.id == 1).first()
    inits = db.query(IniciativaModel).order_by(IniciativaModel.id).all()
    
    inits_list = []
    indice_total = 0.0
    for i in inits:
        pond = round(i.peso_relativo * (i.avance_porcentaje / 100.0), 2)
        indice_total += pond
        inits_list.append({
            "id": i.id,
            "codigo": i.codigo,
            "nombre": i.nombre,
            "descripcion": i.descripcion,
            "eje_engd": i.eje_engd,
            "plazo": i.plazo,
            "peso_relativo": i.peso_relativo,
            "meta_porcentaje": i.meta_porcentaje,
            "avance_porcentaje": i.avance_porcentaje,
            "responsable": i.responsable,
            "puesto": i.puesto,
            "unidad_organica": i.unidad_organica,
            "costo_proyectado": i.costo_proyectado,
            "costo_ejecutado": i.costo_ejecutado,
            "fuente_financiamiento": i.fuente_financiamiento,
            "fecha_inicio": i.fecha_inicio,
            "fecha_fin": i.fecha_fin,
            "indice_ponderado": pond
        })
        
    return {
        "configuracion": {
            "entidad_nombre": cfg.entidad_nombre,
            "siglas": cfg.siglas,
            "ogd_nombre": cfg.ogd_nombre,
            "ogd_puesto": cfg.ogd_puesto,
            "presidente_cgtd": cfg.presidente_cgtd,
            "resolucion_cgtd": cfg.resolucion_cgtd,
            "periodo": cfg.periodo,
            "nivel_objetivo": cfg.nivel_objetivo
        } if cfg else {},
        "iniciativas": inits_list,
        "indice_cumplimiento_total": round(indice_total, 2)
    }

@app.post("/api/plan-accion/config")
def guardar_config_plan(data: ConfigInstitucionalIn, db: Session = Depends(get_db)):
    cfg = db.query(ConfigInstitucionalModel).filter(ConfigInstitucionalModel.id == 1).first()
    if cfg:
        cfg.entidad_nombre = data.entidad_nombre
        cfg.siglas = data.siglas
        cfg.ogd_nombre = data.ogd_nombre
        cfg.ogd_puesto = data.ogd_puesto
        cfg.presidente_cgtd = data.presidente_cgtd
        cfg.resolucion_cgtd = data.resolucion_cgtd
        cfg.periodo = data.periodo
        cfg.nivel_objetivo = data.nivel_objetivo
    else:
        db.add(ConfigInstitucionalModel(id=1, **data.model_dump()))
    db.commit()
    return {"status": "ok", "message": "Configuración institucional guardada"}

@app.post("/api/iniciativas")
def crear_iniciativa(data: IniciativaIn, db: Session = Depends(get_db)):
    init = IniciativaModel(**data.model_dump())
    db.add(init)
    db.commit()
    db.refresh(init)
    return {"status": "ok", "id": init.id, "message": "Iniciativa creada"}

@app.put("/api/iniciativas/{id}")
def actualizar_iniciativa(id: int, data: IniciativaIn, db: Session = Depends(get_db)):
    init = db.query(IniciativaModel).filter(IniciativaModel.id == id).first()
    if not init:
        raise HTTPException(status_code=404, detail="Iniciativa no encontrada")
    for key, val in data.model_dump().items():
        setattr(init, key, val)
    db.commit()
    return {"status": "ok", "message": "Iniciativa actualizada"}

@app.delete("/api/iniciativas/{id}")
def eliminar_iniciativa(id: int, db: Session = Depends(get_db)):
    init = db.query(IniciativaModel).filter(IniciativaModel.id == id).first()
    if init:
        db.delete(init)
        db.commit()
    return {"status": "ok", "message": "Iniciativa eliminada"}

# -------------------------------------------------------------
# EXPORTACIÓN EXCEL Y WORD
# -------------------------------------------------------------
@app.get("/api/exportar/matriz-excel")
def exportar_matriz_excel(db: Session = Depends(get_db)):
    src_file = os.path.join(TEMPLATES_DIR, "Matriz de Evaluación - SISTEMA.xlsx")
    out_file = os.path.join(UPLOADS_DIR, "Matriz_de_Evaluacion_GENERADA.xlsx")
    
    if not os.path.exists(src_file):
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")
    
    wb = openpyxl.load_workbook(src_file)
    ws_eval = wb["Matriz de Evaluación"]
    
    row_mapping = {
        '1.1': 3, '1.2': 4, '1.3': 5, '1.4': 6, '1.5': 7, '1.6': 8, '1.7': 9, '1.8': 10, '1.9': 11, '1.10': 12, '1.11': 13, '1.12': 14,
        '2.1': 16, '2.2': 17, '2.3': 18, '2.4': 19, '2.5': 20,
        '3.1': 22, '3.2': 23, '3.3': 24, '3.4': 25, '3.5': 26, '3.6': 27, '3.7': 28, '3.8': 29, '3.9': 30, '3.10': 31, '3.11': 32,
        '4.1': 34, '4.2': 35, '4.3': 36, '4.4': 37, '4.5': 38, '4.6': 39, '4.7': 40, '4.8': 41, '4.9': 42, '4.10': 43, '4.11': 44, '4.12': 45,
        '5.1': 47, '5.2': 48, '5.3': 49, '5.4': 50, '5.5': 51, '5.6': 52, '5.7': 53, '5.8': 54, '5.9': 55,
        '6.1': 57, '6.2': 58, '6.3': 59, '6.4': 60, '6.5': 61, '6.6': 62, '6.7': 63, '6.8': 64, '6.9': 65,
        '7.1': 67, '7.2': 68, '7.3': 69, '7.4': 70, '7.5': 71, '7.6': 72, '7.7': 73, '7.8': 74, '7.9': 75,
        '8.1': 77, '8.2': 78, '8.3': 79, '8.4': 80, '8.5': 81, '8.6': 82, '8.7': 83, '8.8': 84
    }
    
    respuestas = db.query(RespuestaModel).all()
    for r in respuestas:
        if r.pregunta_codigo in row_mapping:
            t_row = row_mapping[r.pregunta_codigo]
            ws_eval.cell(t_row, 3).value = r.opcion_texto
            ws_eval.cell(t_row, 4).value = r.comentarios
            
    wb.save(out_file)
    return FileResponse(out_file, filename="Matriz_de_Evaluacion_Oficial_PCM.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/api/exportar/evidencias-excel")
def exportar_evidencias_excel(db: Session = Depends(get_db)):
    src_file = os.path.join(TEMPLATES_DIR, "Inventario de evidencias.xlsx")
    out_file = os.path.join(UPLOADS_DIR, "Inventario_de_Evidencias_GENERADO.xlsx")
    
    if not os.path.exists(src_file):
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")
        
    wb = openpyxl.load_workbook(src_file)
    sheet_dim_map = {
        '1': '01 ESTRATEGIA', '2': '02 FINANCIERA', '3': '03 ORGANIZACI\xd3N',
        '4': '04 ARQUITECTURA', '5': '05 CALIDAD', '6': '06 SEGURIDAD',
        '7': '07 OPERACIONES', '8': '08 INTEROPERABILIDAD'
    }
    
    evs = db.query(EvidenciaModel, PreguntaModel.enunciado)\
        .outerjoin(PreguntaModel, EvidenciaModel.pregunta_codigo == PreguntaModel.codigo)\
        .order_by(EvidenciaModel.pregunta_codigo, EvidenciaModel.id).all()
        
    sheet_rows = {s_name: 6 for s_name in wb.sheetnames if s_name.startswith('0')}
    
    for ev, preg_enunciado in evs:
        d_num = ev.pregunta_codigo.split('.')[0] if '.' in ev.pregunta_codigo else '1'
        matched_sheet = None
        for s_key, s_name in sheet_dim_map.items():
            if s_key == d_num and s_name in wb.sheetnames:
                matched_sheet = s_name
                break
        if not matched_sheet:
            for s_name in wb.sheetnames:
                if s_name.startswith(f"0{d_num}"):
                    matched_sheet = s_name
                    break
                    
        if matched_sheet and matched_sheet in wb.sheetnames:
            ws = wb[matched_sheet]
            curr_r = sheet_rows.get(matched_sheet, 6)
            ws.cell(curr_r, 1).value = ev.pregunta_codigo
            ws.cell(curr_r, 2).value = preg_enunciado
            ws.cell(curr_r, 3).value = ev.origen
            ws.cell(curr_r, 4).value = ev.tipo_evidencia
            ws.cell(curr_r, 5).value = ev.nombre_evidencia
            ws.cell(curr_r, 6).value = ev.ruta_o_enlace
            ws.cell(curr_r, 7).value = ev.descripcion
            ws.cell(curr_r, 8).value = ev.responsable
            ws.cell(curr_r, 9).value = ev.area_unidad
            ws.cell(curr_r, 10).value = ev.fecha_registro
            sheet_rows[matched_sheet] = curr_r + 1
            
    wb.save(out_file)
    return FileResponse(out_file, filename="Inventario_de_Evidencias_Oficial_PCM.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/api/exportar/plan-docx")
def exportar_plan_docx(db: Session = Depends(get_db)):
    src_file = os.path.join(TEMPLATES_DIR, "Plan_Accion_GobDatos_PLANTILLA.docx")
    out_file = os.path.join(UPLOADS_DIR, "Plan_de_Accion_Gobierno_Datos_GENERADO.docx")
    
    if not os.path.exists(src_file):
        raise HTTPException(status_code=404, detail="Plantilla Word no encontrada")
        
    doc = docx.Document(src_file)
    cfg = db.query(ConfigInstitucionalModel).filter(ConfigInstitucionalModel.id == 1).first()
    inits = db.query(IniciativaModel).order_by(IniciativaModel.id).all()
    
    entidad_nombre = cfg.entidad_nombre if cfg else "Oficina de Tecnologías de la Información"
    siglas = cfg.siglas if cfg else "OTI"
    ogd_nombre = cfg.ogd_nombre if cfg else "Oficial de Gobierno de Datos"
    ogd_puesto = cfg.ogd_puesto if cfg else "Jefe de Tecnologías"
    resolucion = cfg.resolucion_cgtd if cfg else "Resolución N° 001-2026-CGTD"
    
    replacements = {
        "[NOMBRE COMPLETO DE LA ENTIDAD]": entidad_nombre,
        "[SIGLAS]": siglas,
        "[Apellidos y Nombres]": ogd_nombre,
        "[Puesto]": ogd_puesto,
        "Resolución N.° ______": resolucion,
    }
    
    for p in doc.paragraphs:
        for k, v in replacements.items():
            if k in p.text:
                p.text = p.text.replace(k, v)
                
    if len(doc.tables) > 2:
        t2 = doc.tables[2]
        if len(t2.rows) >= 4:
            t2.rows[1].cells[1].text = entidad_nombre
            t2.rows[2].cells[1].text = siglas
            
    if len(doc.tables) > 3:
        t3 = doc.tables[3]
        if len(t3.rows) >= 3:
            t3.rows[1].cells[1].text = ogd_nombre
            t3.rows[2].cells[1].text = ogd_puesto

    if len(doc.tables) > 9:
        t9 = doc.tables[9]
        dims = db.query(DimensionModel).order_by(DimensionModel.orden).all()
        for idx, d in enumerate(dims):
            if idx + 1 < len(t9.rows):
                pregs = db.query(PreguntaModel).filter(PreguntaModel.dimension_codigo == d.codigo).all()
                preg_codes = [p.codigo for p in pregs]
                resp_rows = db.query(RespuestaModel).filter(RespuestaModel.pregunta_codigo.in_(preg_codes)).all()
                score = sum(r.valor_nivel for r in resp_rows) / len(resp_rows) if resp_rows else 0.0
                t9.rows[idx+1].cells[2].text = f"{score:.2f}"
                t9.rows[idx+1].cells[3].text = "3.00"
                t9.rows[idx+1].cells[4].text = f"{max(0.0, 3.00 - score):.2f}"
                
    if len(doc.tables) > 14 and len(inits) > 0:
        t14 = doc.tables[14]
        for idx, init in enumerate(inits):
            if idx + 1 < len(t14.rows):
                row = t14.rows[idx + 1]
                row.cells[0].text = str(idx + 1)
                row.cells[1].text = init.nombre
                row.cells[2].text = f"{init.meta_porcentaje}%"
                row.cells[3].text = f"{init.peso_relativo}%"
                row.cells[4].text = f"{init.avance_porcentaje}%"
                p_ind = round(init.peso_relativo * (init.avance_porcentaje / 100.0), 2)
                row.cells[5].text = f"{p_ind}%"

    doc.save(out_file)
    return FileResponse(out_file, filename=f"Plan_de_Accion_Gobierno_Datos_{siglas}.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

# Montar estáticos
app.mount("/uploads", StaticFiles(directory=UPLOADS_DIR), name="uploads")

# Montar frontend compilado si existe dist, si no montar static
if os.path.exists(os.path.join(FRONTEND_DIST, "index.html")):
    app.mount("/", StaticFiles(directory=FRONTEND_DIST, html=True), name="frontend_dist")
elif os.path.exists(STATIC_DIR):
    app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
