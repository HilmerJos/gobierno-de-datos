# -*- coding: utf-8 -*-
import os
import json
import shutil
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(os.path.dirname(BASE_DIR))
NUEVO_DIR = os.path.join(PROJECT_ROOT, "MATRIZ DE EVALUACIÓN - NUEVO")
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")

F_ANEXO1 = os.path.join(NUEVO_DIR, "Anexo 1 - MEMGD - Modelo de Evaluacion de Madurez en Gobierno y Gestion de Datos (1).docx")
F_ANEXO2 = os.path.join(NUEVO_DIR, "Anexo 2 - MEMGD - Matriz de Evaluacion (1).xlsx")
F_ANEXO3 = os.path.join(NUEVO_DIR, "Anexo 3 - MEMGD - Inventario de Evidencias (1).xlsx")

CATALOG_BACKEND = os.path.join(BASE_DIR, "catalogo_preguntas_pcm.json")
CATALOG_ROOT = os.path.join(PROJECT_ROOT, "catalogo_preguntas_pcm.json")

print("[MEMGD v1.1] Iniciando extracción de datos desde:", F_ANEXO2)

wb = openpyxl.load_workbook(F_ANEXO2, data_only=True)
ws = wb["Matriz de Evaluación"]

dims_meta = {
    "D1": {
        "id": 1,
        "code": "D1",
        "name": "Estrategia de Gobierno de Datos",
        "description": "Consolida la formulación y aprobación del Plan de Acción de Gobierno de Datos (PAGD), su seguimiento mediante indicadores, su articulación con los instrumentos de planeamiento institucional, y la gestión de los datos como activo estratégico en la operación diaria."
    },
    "D2": {
        "id": 2,
        "code": "D2",
        "name": "Viabilidad Financiera",
        "description": "Se enfoca en la justificación económica y la asignación presupuestaria necesaria para la sostenibilidad de las iniciativas de gobierno de datos, incluyendo la identificación de casos de uso de valor."
    },
    "D3": {
        "id": 3,
        "code": "D3",
        "name": "Organización y Cultura de Datos",
        "description": "Se centra en la estructura de gobernanza, los roles y responsabilidades, las políticas institucionales y los aspectos culturales (ética, capacitación, comunicación) necesarios para gobernar el dato."
    },
    "D4": {
        "id": 4,
        "code": "D4",
        "name": "Arquitectura de Datos",
        "description": "Es la base conceptual y de diseño del dato, abarcando la identificación de datos maestros, el Modelo de Referencia de Datos (MRD), la estandarización semántica y la gestión de metadatos."
    },
    "D5": {
        "id": 5,
        "code": "D5",
        "name": "Calidad de Datos",
        "description": "Cubre la gestión integral del Programa de Calidad de Datos (PCD), desde el marco normativo hasta la medición, monitoreo, corrección de problemas y valoración económica del dato."
    },
    "D6": {
        "id": 6,
        "code": "D6",
        "name": "Seguridad de Datos",
        "description": "Abarca la seguridad de la información, la gestión de riesgos, la protección de datos personales, el cumplimiento normativo obligatorio y los riesgos de sistemas basados en Inteligencia Artificial."
    },
    "D7": {
        "id": 7,
        "code": "D7",
        "name": "Infraestructura Tecnológica y Operaciones de Datos",
        "description": "Se centra en la tecnología, infraestructura, continuidad operativa, integración de sistemas, analítica avanzada, gobernanza de servicios en la nube y la gestión operativa del ciclo de vida del dato."
    },
    "D8": {
        "id": 8,
        "code": "D8",
        "name": "Interoperabilidad y Datos Abiertos",
        "description": "Agrupa los mecanismos para el intercambio, la publicación y el consumo de datos de forma externa, mediante las plataformas nacionales de interoperabilidad y datos abiertos."
    }
}

level_labels = {
    0: "No iniciado",
    1: "Inicial",
    2: "En desarrollo",
    3: "Definido",
    4: "Logrado",
    5: "Optimizado"
}

catalog = []
dim_map = {}
for code, meta in dims_meta.items():
    d_obj = {
        "id": meta["id"],
        "code": code,
        "name": meta["name"],
        "description": meta["description"],
        "questions": []
    }
    catalog.append(d_obj)
    dim_map[code] = d_obj

q_id = 1
for r in range(2, ws.max_row + 1):
    q_code = ws.cell(r, 2).value
    if not q_code:
        continue
    q_code = str(q_code).strip()
    d_num = q_code.split(".")[0]
    d_code = f"D{d_num}"
    
    title = str(ws.cell(r, 3).value or "").strip()
    enunciado = str(ws.cell(r, 4).value or "").strip()
    admite_na = str(ws.cell(r, 11).value or "No").strip()
    explicacion = str(ws.cell(r, 13).value or "").strip()
    resultados = str(ws.cell(r, 14).value or "").strip()
    
    options = []
    for lvl in range(6):
        desc = ws.cell(r, 5 + lvl).value
        options.append({
            "nivel": lvl,
            "etiqueta": level_labels[lvl],
            "descripcion": str(desc or "").strip()
        })
        
    q_obj = {
        "id": q_id,
        "code": q_code,
        "dimension_codigo": d_code,
        "titulo": title,
        "enunciado": enunciado,
        "explicacion": explicacion,
        "resultados_esperados": resultados,
        "admite_na": admite_na,
        "options": options
    }
    
    if d_code in dim_map:
        dim_map[d_code]["questions"].append(q_obj)
    q_id += 1

total_q = sum(len(d["questions"]) for d in catalog)
print(f"[MEMGD v1.1] Catálogo generado con éxito: {len(catalog)} dimensiones, {total_q} preguntas.")
for d in catalog:
    print(f"  - {d['code']}: {d['name']} -> {len(d['questions'])} preguntas")

# Guardar catálogo JSON en backend y raíz
with open(CATALOG_BACKEND, "w", encoding="utf-8") as f:
    json.dump(catalog, f, ensure_ascii=False, indent=2)
print("[OK] Guardado en:", CATALOG_BACKEND)

with open(CATALOG_ROOT, "w", encoding="utf-8") as f:
    json.dump(catalog, f, ensure_ascii=False, indent=2)
print("[OK] Guardado en:", CATALOG_ROOT)

# Actualizar plantillas en backend/templates
os.makedirs(TEMPLATES_DIR, exist_ok=True)
dest_matriz = os.path.join(TEMPLATES_DIR, "Matriz de Evaluación - SISTEMA.xlsx")
dest_evidencias = os.path.join(TEMPLATES_DIR, "Inventario de evidencias.xlsx")

shutil.copy2(F_ANEXO2, dest_matriz)
print("[OK] Plantilla Matriz actualizada:", dest_matriz)

shutil.copy2(F_ANEXO3, dest_evidencias)
print("[OK] Plantilla Inventario actualizada:", dest_evidencias)

print("[MEMGD v1.1] Proceso de sincronización de catálogo y plantillas finalizado exitosamente.")
